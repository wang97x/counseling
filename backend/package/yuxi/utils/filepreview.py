"""与存储和 HTTP 无关的文件预览渲染原语。"""

from __future__ import annotations

import asyncio
import base64
import html
import mimetypes
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePosixPath

MAX_BINARY_PREVIEW_SIZE_BYTES = 30 * 1024 * 1024
MAX_TEXT_PREVIEW_CHARS = 250_000
MAX_OFFICE_ARCHIVE_ENTRIES = 1_500
MAX_OFFICE_ARCHIVE_ENTRY_BYTES = 16 * 1024 * 1024
MAX_OFFICE_ARCHIVE_UNCOMPRESSED_BYTES = 40 * 1024 * 1024
MAX_OFFICE_COMPRESSION_RATIO = 200
MAX_OFFICE_EMBEDDED_IMAGE_BYTES = 4 * 1024 * 1024
MAX_OFFICE_HTML_CHARS = 6 * 1024 * 1024

_MARKDOWN_EXTENSIONS = frozenset({".md", ".markdown", ".mdx"})
_PDF_EXTENSIONS = frozenset({".pdf"})
_HTML_EXTENSIONS = frozenset({".html", ".htm"})
_OFFICE_HTML_PREVIEW_EXTENSIONS = frozenset({".docx", ".pptx"})
_OFFICE_MEDIA_TYPES = {
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
_TEXT_EXTENSIONS = frozenset(
    {
        ".txt",
        ".text",
        ".log",
        ".json",
        ".jsonl",
        ".yaml",
        ".yml",
        ".toml",
        ".ini",
        ".cfg",
        ".conf",
        ".csv",
        ".tsv",
        ".py",
        ".js",
        ".ts",
        ".jsx",
        ".tsx",
        ".vue",
        ".css",
        ".less",
        ".scss",
        ".xml",
        ".sql",
        ".sh",
        ".bash",
        ".zsh",
        ".fish",
        ".env",
        ".dockerfile",
        ".gitignore",
        ".weather",
    }
)
_IMAGE_EXTENSIONS = frozenset({".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".svg"})
_BINARY_SIGNATURES = (
    b"\x7fELF",
    b"MZ",
    b"%PDF-",
    b"PK\x03\x04",
    b"PK\x05\x06",
    b"PK\x07\x08",
    b"\x89PNG\r\n\x1a\n",
    b"\xff\xd8\xff",
    b"GIF87a",
    b"GIF89a",
    b"RIFF",
)


class OfficePreviewConversionError(RuntimeError):
    """Office 文件无法生成结构化预览。"""


@dataclass(frozen=True, slots=True)
class PreviewResult:
    """与存储和 HTTP 无关的文件预览结果。"""

    content: str | bytes | None
    preview_type: str
    supported: bool
    media_type: str | None = None
    filename: str | None = None
    message: str | None = None
    truncated: bool = False
    limit: int | None = None

    def payload(self) -> dict:
        """返回文本或不支持预览使用的数据。"""
        return {
            "content": self.content if isinstance(self.content, str) else None,
            "preview_type": self.preview_type,
            "supported": self.supported,
            "message": self.message,
            "truncated": self.truncated,
            "limit": self.limit,
        }


def is_office_html_preview_file(path: str) -> bool:
    """判断文件是否支持轻量 Office HTML 预览。"""
    return PurePosixPath(path).suffix.lower() in _OFFICE_HTML_PREVIEW_EXTENSIONS


def preview_too_large() -> PreviewResult:
    return PreviewResult(
        content=None,
        preview_type="unsupported",
        supported=False,
        message="文件过大，当前仅支持 30 MB 以内的文件预览",
        limit=MAX_BINARY_PREVIEW_SIZE_BYTES,
    )


def _html_document(title: str, body: str) -> str:
    """生成不包含脚本和外部资源的独立预览页面。"""
    safe_title = html.escape(title)
    document = f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><title>{safe_title}</title>
<style>
body{{margin:0;background:#f3f4f6;color:#1f2937;font:15px/1.65 system-ui,sans-serif}}
main{{box-sizing:border-box;max-width:960px;min-height:100vh;margin:0 auto;padding:40px 48px;background:#fff}}
h1,h2,h3,h4,h5,h6{{line-height:1.3}} table{{width:100%;border-collapse:collapse;margin:16px 0}}
th,td{{border:1px solid #d1d5db;padding:6px 9px;text-align:left;vertical-align:top}}
.slide{{margin:0 0 24px;padding:24px;border:1px solid #d1d5db;border-radius:8px}}
.slide-number{{color:#6b7280;font-size:12px}} .empty{{color:#6b7280}}
.office-images{{display:flex;flex-wrap:wrap;gap:12px;margin-top:16px}}
.office-images img{{max-width:100%;max-height:560px;object-fit:contain}}
</style></head><body><main>{body}</main></body></html>"""
    if len(document) > MAX_OFFICE_HTML_CHARS:
        raise OfficePreviewConversionError("Office 预览内容超过 HTML 大小限制")
    return document


def _validate_office_archive(content: bytes) -> None:
    """在进程内解析前约束 OOXML ZIP 的解压资源。"""
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries = [entry for entry in archive.infolist() if not entry.is_dir()]
    except zipfile.BadZipFile as exc:
        raise OfficePreviewConversionError("Office 文件不是有效的 OOXML 压缩包") from exc

    if len(entries) > MAX_OFFICE_ARCHIVE_ENTRIES:
        raise OfficePreviewConversionError("Office 文件包含过多压缩条目")

    total_size = 0
    for entry in entries:
        if entry.flag_bits & 0x1:
            raise OfficePreviewConversionError("Office 文件包含不支持的加密条目")
        if entry.file_size > MAX_OFFICE_ARCHIVE_ENTRY_BYTES:
            raise OfficePreviewConversionError("Office 文件包含过大的压缩条目")
        if entry.file_size and (
            entry.compress_size == 0 or entry.file_size / entry.compress_size > MAX_OFFICE_COMPRESSION_RATIO
        ):
            raise OfficePreviewConversionError("Office 文件压缩比超过安全限制")
        total_size += entry.file_size
        if total_size > MAX_OFFICE_ARCHIVE_UNCOMPRESSED_BYTES:
            raise OfficePreviewConversionError("Office 文件解压后超过大小限制")


def _table_html(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    rendered_rows = []
    for index, row in enumerate(rows):
        tag = "th" if index == 0 else "td"
        rendered_cells = "".join(f"<{tag}>{html.escape(cell)}</{tag}>" for cell in row)
        rendered_rows.append(f"<tr>{rendered_cells}</tr>")
    return f"<table><tbody>{''.join(rendered_rows)}</tbody></table>"


def _image_html(content_type: str, blob: bytes, alt: str) -> str:
    if content_type not in {"image/png", "image/jpeg", "image/gif", "image/webp"}:
        return ""
    encoded = base64.b64encode(blob).decode("ascii")
    return f'<img src="data:{content_type};base64,{encoded}" alt="{html.escape(alt, quote=True)}">'


def _convert_docx_to_html(content: bytes, title: str) -> str:
    from docx import Document

    document = Document(BytesIO(content))
    blocks: list[str] = []
    for block in document.iter_inner_content():
        if hasattr(block, "rows"):
            rows = [[cell.text.strip() for cell in row.cells] for row in block.rows]
            blocks.append(_table_html(rows))
            continue
        text = block.text.strip()
        if not text:
            continue
        style_name = str(block.style.name or "").lower()
        if style_name.startswith("heading"):
            level_text = style_name.removeprefix("heading").strip()
            level = min(max(int(level_text), 1), 6) if level_text.isdigit() else 2
            blocks.append(f"<h{level}>{html.escape(text)}</h{level}>")
        else:
            blocks.append(f"<p>{html.escape(text)}</p>")
    images = []
    image_bytes = 0
    for part in document.part.related_parts.values():
        content_type = str(getattr(part, "content_type", ""))
        blob = getattr(part, "blob", None)
        if content_type.startswith("image/") and isinstance(blob, bytes):
            image_bytes += len(blob)
            if image_bytes > MAX_OFFICE_EMBEDDED_IMAGE_BYTES:
                raise OfficePreviewConversionError("Office 文件内嵌图片超过预览大小限制")
            rendered = _image_html(content_type, blob, "文档图片")
            if rendered:
                images.append(rendered)
    if images:
        blocks.append(f'<div class="office-images">{"".join(images)}</div>')
    body = "".join(blocks) or '<p class="empty">文档中没有可预览的文本或表格。</p>'
    return _html_document(title, body)


def _convert_pptx_to_html(content: bytes, title: str) -> str:
    from pptx import Presentation

    presentation = Presentation(BytesIO(content))
    slides: list[str] = []
    image_bytes = 0
    for slide_number, slide in enumerate(presentation.slides, start=1):
        blocks = [f'<div class="slide-number">第 {slide_number} 页</div>']
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                blocks.append(_table_html(rows))
                continue
            text = str(getattr(shape, "text", "")).strip()
            if text:
                blocks.append(f"<p>{html.escape(text).replace(chr(10), '<br>')}</p>")
            image = getattr(shape, "image", None)
            if image is not None:
                blob = image.blob
                image_bytes += len(blob)
                if image_bytes > MAX_OFFICE_EMBEDDED_IMAGE_BYTES:
                    raise OfficePreviewConversionError("Office 文件内嵌图片超过预览大小限制")
                rendered = _image_html(str(image.content_type), blob, "幻灯片图片")
                if rendered:
                    blocks.append(f'<div class="office-images">{rendered}</div>')
        slides.append(f'<section class="slide">{"".join(blocks)}</section>')
    body = "".join(slides) or '<p class="empty">演示文稿中没有可预览的文本或表格。</p>'
    return _html_document(title, body)


def _convert_office_to_html_sync(filename: str, content: bytes) -> str:
    suffix = PurePosixPath(filename).suffix.lower()
    if suffix not in _OFFICE_HTML_PREVIEW_EXTENSIONS:
        raise OfficePreviewConversionError("当前文件类型不支持结构化预览")
    try:
        _validate_office_archive(content)
        if suffix == ".docx":
            return _convert_docx_to_html(content, filename)
        return _convert_pptx_to_html(content, filename)
    except OfficePreviewConversionError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise OfficePreviewConversionError(f"Office 文件预览失败: {exc}") from exc


async def convert_office_to_html(filename: str, content: bytes) -> str:
    """把 DOCX/PPTX 字节转换为轻量结构化 HTML。"""
    return await asyncio.to_thread(_convert_office_to_html_sync, filename, content)


def detect_preview_type(path: str, raw_content: bytes) -> tuple[str, bool, str | None]:
    suffix = PurePosixPath(path).suffix.lower()
    mime_type, _encoding = mimetypes.guess_type(path)
    head = raw_content[:1024]

    if suffix in _IMAGE_EXTENSIONS or (mime_type and mime_type.startswith("image/")):
        return "image", True, None
    if suffix in _PDF_EXTENSIONS or mime_type == "application/pdf" or head.startswith(b"%PDF-"):
        return "pdf", True, None
    if suffix in _MARKDOWN_EXTENSIONS:
        return "markdown", True, None
    if suffix in _HTML_EXTENSIONS:
        return "html", True, None
    if suffix in _TEXT_EXTENSIONS:
        return "text", True, None
    if b"\x00" in head:
        return "unsupported", False, "当前文件是二进制文件，暂不支持预览"
    if any(head.startswith(signature) for signature in _BINARY_SIGNATURES):
        if head.startswith(b"RIFF") and b"WEBP" in head[:16]:
            return "image", True, None
        return "unsupported", False, "当前文件格式暂不支持预览"
    if mime_type:
        if mime_type.startswith("text/"):
            return "text", True, None
        if mime_type in {"application/json", "application/xml", "application/javascript"}:
            return "text", True, None
        if mime_type.startswith("application/"):
            return "unsupported", False, "当前文件格式暂不支持预览"
    if not raw_content:
        return "text", True, None
    try:
        raw_content.decode("utf-8")
        return "text", True, None
    except UnicodeDecodeError:
        return "unsupported", False, "当前文件不是可读文本，暂不支持预览"


def render_preview(path: str, raw_content: bytes) -> PreviewResult:
    """把文件字节渲染为中立 Preview 结果。"""
    preview_type, supported, message = detect_preview_type(path, raw_content)
    if preview_type in {"image", "pdf"}:
        return PreviewResult(
            content=raw_content,
            preview_type=preview_type,
            supported=True,
            media_type=detect_media_type(path, raw_content),
            filename=PurePosixPath(path).name or "preview",
        )
    if not supported:
        return PreviewResult(
            content=None,
            preview_type=preview_type,
            supported=supported,
            message=message,
        )

    try:
        content = raw_content.decode("utf-8")
    except UnicodeDecodeError:
        return PreviewResult(
            content=None,
            preview_type="unsupported",
            supported=False,
            message="当前文件不是 UTF-8 文本，暂不支持预览",
        )

    truncated = len(content) > MAX_TEXT_PREVIEW_CHARS
    if truncated:
        content = content[:MAX_TEXT_PREVIEW_CHARS]
    return PreviewResult(
        content=content,
        preview_type=preview_type,
        supported=True,
        message=message,
        truncated=truncated,
        limit=MAX_TEXT_PREVIEW_CHARS,
    )


def detect_media_type(path: str, raw_content: bytes | None = None) -> str:
    """优先按文件签名识别响应媒体类型。"""
    head = (raw_content or b"")[:512]
    if head.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if head.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if head.startswith(b"GIF87a") or head.startswith(b"GIF89a"):
        return "image/gif"
    if head.startswith(b"RIFF") and b"WEBP" in head[:16]:
        return "image/webp"
    if head.startswith(b"BM"):
        return "image/bmp"
    if head.startswith(b"%PDF-"):
        return "application/pdf"

    stripped_head = head.lstrip()
    if stripped_head.startswith(b"<svg") or stripped_head.startswith(b"<?xml"):
        suffix = PurePosixPath(path).suffix.lower()
        if suffix == ".svg" or b"<svg" in stripped_head[:256]:
            return "image/svg+xml"

    suffix = PurePosixPath(path).suffix.lower()
    if suffix in _OFFICE_MEDIA_TYPES:
        return _OFFICE_MEDIA_TYPES[suffix]
    return mimetypes.guess_type(path)[0] or "application/octet-stream"
