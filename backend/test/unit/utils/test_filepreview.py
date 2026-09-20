from __future__ import annotations

import base64
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
import yuxi.utils.filepreview as filepreview
from docx import Document
from pptx import Presentation

from yuxi.utils.filepreview import (
    MAX_TEXT_PREVIEW_CHARS,
    OfficePreviewConversionError,
    convert_office_to_html,
    detect_preview_type,
    is_office_html_preview_file,
    render_preview,
)


def _build_docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


_ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _build_docx_with_image() -> bytes:
    document = Document()
    document.add_picture(BytesIO(_ONE_PIXEL_PNG))
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_pptx_with_image() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(BytesIO(_ONE_PIXEL_PNG), 0, 0)
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _build_zip(entries: dict[str, bytes]) -> bytes:
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    return buffer.getvalue()


@pytest.mark.parametrize(
    ("renderer", "expected_third"),
    [
        (detect_preview_type, "当前文件是二进制文件，暂不支持预览"),
        (render_preview, None),
    ],
)
def test_docx_is_not_treated_as_markdown_preview(renderer, expected_third):
    result = renderer("demo.docx", _build_docx_bytes("Docx preview"))

    if isinstance(result, tuple):
        preview_type, supported, third = result
    else:
        preview_type, supported, third = result.preview_type, result.supported, result.content

    assert preview_type == "unsupported"
    assert supported is False
    assert third == expected_third


def test_render_preview_truncates_long_markdown():
    result = render_preview("note.md", ("x" * (MAX_TEXT_PREVIEW_CHARS + 1)).encode("utf-8"))

    assert result.preview_type == "markdown"
    assert result.supported is True
    assert result.truncated is True
    assert result.limit == MAX_TEXT_PREVIEW_CHARS
    assert len(result.content) == MAX_TEXT_PREVIEW_CHARS


def test_render_preview_returns_complete_binary_result_from_signature():
    content = b"%PDF-1.4\npreview"

    result = render_preview("report.bin", content)

    assert result.content == content
    assert result.preview_type == "pdf"
    assert result.supported is True
    assert result.media_type == "application/pdf"
    assert result.filename == "report.bin"


def test_render_preview_keeps_unsupported_binary_content_hidden():
    result = render_preview("archive.bin", b"\x00binary")

    assert result.content is None
    assert result.preview_type == "unsupported"
    assert result.supported is False


def test_office_html_preview_scope_only_includes_docx_and_pptx():
    assert is_office_html_preview_file("demo.docx") is True
    assert is_office_html_preview_file("demo.pptx") is True
    assert is_office_html_preview_file("demo.xlsx") is False
    assert is_office_html_preview_file("demo.doc") is False
    assert is_office_html_preview_file("demo.ppt") is False


@pytest.mark.asyncio
async def test_docx_preview_returns_escaped_structured_html():
    document = Document()
    document.add_heading("标题 <script>", level=1)
    document.add_paragraph("正文 & 内容")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(1, 0).text = "<img>"
    buffer = BytesIO()
    document.save(buffer)

    result = await convert_office_to_html("demo.docx", buffer.getvalue())

    assert "<h1>标题 &lt;script&gt;</h1>" in result
    assert "正文 &amp; 内容" in result
    assert "&lt;img&gt;" in result
    assert "<script>" not in result


@pytest.mark.asyncio
async def test_pptx_preview_keeps_slide_boundaries_and_escapes_text():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "第一页 <title>"
    slide.placeholders[1].text = "正文 & 内容"
    buffer = BytesIO()
    presentation.save(buffer)

    result = await convert_office_to_html("slides.pptx", buffer.getvalue())

    assert "第 1 页" in result
    assert "第一页 &lt;title&gt;" in result
    assert "正文 &amp; 内容" in result


@pytest.mark.asyncio
async def test_office_html_preview_rejects_corrupt_zip():
    with pytest.raises(OfficePreviewConversionError, match="不是有效的 OOXML"):
        await convert_office_to_html("demo.docx", b"not-a-docx")


@pytest.mark.asyncio
async def test_office_html_preview_rejects_high_compression_archive():
    buffer = BytesIO()
    with ZipFile(buffer, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"0" * (1024 * 1024))

    with pytest.raises(OfficePreviewConversionError, match="压缩比超过安全限制"):
        await convert_office_to_html("bomb.docx", buffer.getvalue())


@pytest.mark.asyncio
async def test_office_html_preview_rejects_excessive_entry_count(monkeypatch):
    monkeypatch.setattr(filepreview, "MAX_OFFICE_ARCHIVE_ENTRIES", 1)

    with pytest.raises(OfficePreviewConversionError, match="过多压缩条目"):
        await convert_office_to_html(
            "many.docx",
            _build_zip({"[Content_Types].xml": b"types", "word/document.xml": b"document"}),
        )


@pytest.mark.asyncio
async def test_office_html_preview_rejects_encrypted_entry():
    content = bytearray(_build_zip({"word/document.xml": b"document"}))
    central_header = content.index(b"PK\x01\x02")
    flag_offset = central_header + 8
    flags = int.from_bytes(content[flag_offset : flag_offset + 2], "little") | 0x1
    content[flag_offset : flag_offset + 2] = flags.to_bytes(2, "little")

    with pytest.raises(OfficePreviewConversionError, match="加密条目"):
        await convert_office_to_html("encrypted.docx", bytes(content))


@pytest.mark.asyncio
async def test_office_html_preview_rejects_oversized_entry(monkeypatch):
    monkeypatch.setattr(filepreview, "MAX_OFFICE_ARCHIVE_ENTRY_BYTES", 3)

    with pytest.raises(OfficePreviewConversionError, match="过大的压缩条目"):
        await convert_office_to_html("large-entry.docx", _build_zip({"word/document.xml": b"1234"}))


@pytest.mark.asyncio
async def test_office_html_preview_rejects_oversized_uncompressed_total(monkeypatch):
    monkeypatch.setattr(filepreview, "MAX_OFFICE_ARCHIVE_UNCOMPRESSED_BYTES", 5)

    with pytest.raises(OfficePreviewConversionError, match="解压后超过大小限制"):
        await convert_office_to_html(
            "large-total.docx",
            _build_zip({"[Content_Types].xml": b"123", "word/document.xml": b"456"}),
        )


@pytest.mark.parametrize(
    ("filename", "builder"),
    [("image.docx", _build_docx_with_image), ("image.pptx", _build_pptx_with_image)],
)
@pytest.mark.asyncio
async def test_office_html_preview_rejects_oversized_embedded_images(monkeypatch, filename, builder):
    monkeypatch.setattr(filepreview, "MAX_OFFICE_EMBEDDED_IMAGE_BYTES", 1)

    with pytest.raises(OfficePreviewConversionError, match="内嵌图片超过预览大小限制"):
        await convert_office_to_html(filename, builder())


@pytest.mark.asyncio
async def test_office_html_preview_rejects_oversized_rendered_html(monkeypatch):
    monkeypatch.setattr(filepreview, "MAX_OFFICE_HTML_CHARS", 64)

    with pytest.raises(OfficePreviewConversionError, match="HTML 大小限制"):
        await convert_office_to_html("demo.docx", _build_docx_bytes("preview"))
